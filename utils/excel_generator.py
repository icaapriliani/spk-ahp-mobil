import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_ahp_excel(kriteria, alternatif, matriks_kriteria, matriks_alternatif_dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Perhitungan AHP"
    
    # Styles
    font_bold = Font(bold=True)
    fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_input = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    fill_result = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    n_k = len(kriteria)
    n_a = len(alternatif)
    
    current_row = 1
    ws.cell(row=current_row, column=1, value="PERHITUNGAN SISTEM PENDUKUNG KEPUTUSAN (AHP)").font = Font(bold=True, size=14)
    current_row += 2
    
    # --- 1. DATA AWAL ---
    ws.cell(row=current_row, column=1, value="1. DATA AWAL").font = font_bold
    current_row += 1
    
    ws.cell(row=current_row, column=1, value="Kriteria").font = font_bold
    ws.cell(row=current_row, column=2, value="Alternatif").font = font_bold
    
    max_len = max(n_k, n_a)
    for i in range(max_len):
        if i < n_k:
            ws.cell(row=current_row + 1 + i, column=1, value=kriteria[i])
        if i < n_a:
            ws.cell(row=current_row + 1 + i, column=2, value=alternatif[i])
            
    current_row += max_len + 2
    
    # Fungsi pembantu untuk membuat tabel dengan border
    def style_range(ws, start_row, start_col, end_row, end_col, fill=None, bold_headers=False):
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border_thin
                if fill:
                    cell.fill = fill
                if bold_headers and (r == start_row or c == start_col):
                    cell.font = font_bold

    # --- 2. PERBANDINGAN KRITERIA ---
    ws.cell(row=current_row, column=1, value="2. MATRIKS PERBANDINGAN KRITERIA").font = font_bold
    current_row += 1
    
    # Header Kriteria
    start_row_mk = current_row
    for i, k in enumerate(kriteria):
        ws.cell(row=current_row, column=i+2, value=k).font = font_bold
        ws.cell(row=current_row+i+1, column=1, value=k).font = font_bold
        
    # Isi Matriks Input (Warna kuning untuk input manual/asli)
    for i in range(n_k):
        for j in range(n_k):
            cell = ws.cell(row=current_row+1+i, column=j+2)
            if i == j:
                cell.value = 1
            elif i < j:
                val = matriks_kriteria.iloc[i, j]
                cell.value = val
                cell.fill = fill_input
            else:
                # Rumus kebalikan: =1/Cell_Atas
                col_letter = get_column_letter(i+2)
                row_num = current_row+1+j
                cell.value = f"=1/{col_letter}{row_num}"
                
    style_range(ws, current_row, 1, current_row+n_k, n_k+1)
    
    # Baris Jumlah Kolom
    current_row += n_k + 1
    ws.cell(row=current_row, column=1, value="Jumlah").font = font_bold
    for j in range(n_k):
        col_letter = get_column_letter(j+2)
        ws.cell(row=current_row, column=j+2, value=f"=SUM({col_letter}{start_row_mk+1}:{col_letter}{current_row-1})")
    style_range(ws, current_row, 1, current_row, n_k+1, fill=fill_header)
    
    current_row += 2
    
    # Normalisasi Matriks Kriteria
    ws.cell(row=current_row, column=1, value="NORMALISASI MATRIKS KRITERIA").font = font_bold
    current_row += 1
    start_row_norm_mk = current_row
    
    for i, k in enumerate(kriteria):
        ws.cell(row=current_row, column=i+2, value=k).font = font_bold
        ws.cell(row=current_row+i+1, column=1, value=k).font = font_bold
        
    ws.cell(row=current_row, column=n_k+2, value="Bobot (Eigen Vector)").font = font_bold
    
    for i in range(n_k):
        for j in range(n_k):
            col_letter = get_column_letter(j+2)
            val_cell = f"{col_letter}{start_row_mk + 1 + i}"
            sum_cell = f"{col_letter}{start_row_mk + n_k + 1}"
            ws.cell(row=current_row+1+i, column=j+2, value=f"={val_cell}/{sum_cell}")
            
        # Bobot = Rata-rata baris
        start_col = get_column_letter(2)
        end_col = get_column_letter(n_k+1)
        ws.cell(row=current_row+1+i, column=n_k+2, value=f"=AVERAGE({start_col}{current_row+1+i}:{end_col}{current_row+1+i})").fill = fill_result
        
    style_range(ws, current_row, 1, current_row+n_k, n_k+2)
    bobot_kriteria_start_row = current_row + 1 # Akan digunakan nanti
    current_row += n_k + 2
    
    # Uji Konsistensi (CR)
    ws.cell(row=current_row, column=1, value="UJI KONSISTENSI (CR)").font = font_bold
    current_row += 1
    
    ws.cell(row=current_row, column=1, value="Kriteria").font = font_bold
    ws.cell(row=current_row, column=2, value="Weighted Sum").font = font_bold
    ws.cell(row=current_row, column=3, value="Rasio Konsistensi").font = font_bold
    
    for i, k in enumerate(kriteria):
        ws.cell(row=current_row+1+i, column=1, value=k)
        # BUKAN SUMPRODUCT KARENA DIMENSI BERBEDA (Horizontal vs Vertikal di Excel menghasilkan #VALUE!)
        # Pakai perkalian manual: =B7*$D$15 + C7*$D$16
        formula_parts = []
        for j in range(n_k):
            cell_asli = f"{get_column_letter(j+2)}{start_row_mk+1+i}"
            cell_bobot = f"${get_column_letter(n_k+2)}${bobot_kriteria_start_row+j}"
            formula_parts.append(f"{cell_asli}*{cell_bobot}")
            
        ws.cell(row=current_row+1+i, column=2, value="=" + "+".join(formula_parts))
        
        # Rasio Konsistensi = Weighted Sum / Bobot
        ws.cell(row=current_row+1+i, column=3, value=f"=B{current_row+1+i}/{get_column_letter(n_k+2)}{bobot_kriteria_start_row+i}")
        
    style_range(ws, current_row, 1, current_row+n_k, 3)
    current_row += n_k + 1
    
    ws.cell(row=current_row, column=2, value="Lambda Max").font = font_bold
    ws.cell(row=current_row, column=3, value=f"=AVERAGE(C{current_row-n_k}:C{current_row-1})")
    
    ws.cell(row=current_row+1, column=2, value="Consistency Index (CI)").font = font_bold
    ws.cell(row=current_row+1, column=3, value=f"=(C{current_row}-{n_k})/({n_k}-1)")
    
    ri_dict = {1: 0.0, 2: 0.0, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24, 7: 1.32, 8: 1.41, 9: 1.45, 10: 1.49}
    ri = ri_dict.get(n_k, 1.49)
    ws.cell(row=current_row+2, column=2, value="Random Index (RI)").font = font_bold
    ws.cell(row=current_row+2, column=3, value=ri)
    
    ws.cell(row=current_row+3, column=2, value="Consistency Ratio (CR)").font = font_bold
    # Mencegah error #DIV/0! jika n < 3 (RI = 0)
    ws.cell(row=current_row+3, column=3, value=f"=IF(C{current_row+2}=0, 0, C{current_row+1}/C{current_row+2})").fill = fill_result
    
    current_row += 5
    
    # --- 3. PERBANDINGAN ALTERNATIF ---
    ws.cell(row=current_row, column=1, value="3. MATRIKS PERBANDINGAN ALTERNATIF (Terhadap tiap Kriteria)").font = font_bold
    current_row += 1
    
    bobot_alternatif_refs = {} # Menyimpan referensi cell bobot alternatif untuk rekap
    
    for k_idx, k in enumerate(kriteria):
        ws.cell(row=current_row, column=1, value=f"Kriteria: {k}").font = font_bold
        current_row += 1
        
        start_row_ma = current_row
        for i, a in enumerate(alternatif):
            ws.cell(row=current_row, column=i+2, value=a).font = font_bold
            ws.cell(row=current_row+i+1, column=1, value=a).font = font_bold
            
        matriks_alt = matriks_alternatif_dict.get(k)
        
        for i in range(n_a):
            for j in range(n_a):
                cell = ws.cell(row=current_row+1+i, column=j+2)
                if i == j:
                    cell.value = 1
                elif i < j:
                    val = matriks_alt.iloc[i, j] if matriks_alt is not None else 1
                    cell.value = val
                    cell.fill = fill_input
                else:
                    col_letter = get_column_letter(i+2)
                    row_num = current_row+1+j
                    cell.value = f"=1/{col_letter}{row_num}"
                    
        style_range(ws, current_row, 1, current_row+n_a, n_a+1)
        
        current_row += n_a + 1
        ws.cell(row=current_row, column=1, value="Jumlah").font = font_bold
        for j in range(n_a):
            col_letter = get_column_letter(j+2)
            ws.cell(row=current_row, column=j+2, value=f"=SUM({col_letter}{start_row_ma+1}:{col_letter}{current_row-1})")
        style_range(ws, current_row, 1, current_row, n_a+1, fill=fill_header)
        
        current_row += 2
        
        # Normalisasi
        ws.cell(row=current_row, column=1, value=f"Normalisasi Kriteria: {k}").font = font_bold
        current_row += 1
        
        for i, a in enumerate(alternatif):
            ws.cell(row=current_row, column=i+2, value=a).font = font_bold
            ws.cell(row=current_row+i+1, column=1, value=a).font = font_bold
            
        ws.cell(row=current_row, column=n_a+2, value="Bobot (Eigen Vector)").font = font_bold
        
        start_row_bobot_alt = current_row + 1
        for i in range(n_a):
            for j in range(n_a):
                col_letter = get_column_letter(j+2)
                val_cell = f"{col_letter}{start_row_ma + 1 + i}"
                sum_cell = f"{col_letter}{start_row_ma + n_a + 1}"
                ws.cell(row=current_row+1+i, column=j+2, value=f"={val_cell}/{sum_cell}")
                
            start_col = get_column_letter(2)
            end_col = get_column_letter(n_a+1)
            cell_bobot = ws.cell(row=current_row+1+i, column=n_a+2, value=f"=AVERAGE({start_col}{current_row+1+i}:{end_col}{current_row+1+i})")
            cell_bobot.fill = fill_result
            
        style_range(ws, current_row, 1, current_row+n_a, n_a+2)
        
        bobot_alternatif_refs[k] = (start_row_bobot_alt, start_row_bobot_alt + n_a - 1, get_column_letter(n_a+2))
        
        current_row += n_a + 3

    # --- 4. HASIL AKHIR ---
    ws.cell(row=current_row, column=1, value="4. HASIL AKHIR").font = font_bold
    current_row += 1
    
    # Headers
    ws.cell(row=current_row, column=1, value="Alternatif").font = font_bold
    for i, k in enumerate(kriteria):
        ws.cell(row=current_row, column=i+2, value=k).font = font_bold
        # Baris Bobot Kriteria
        col_bobot_kriteria = get_column_letter(n_k+2)
        ws.cell(row=current_row+1, column=i+2, value=f"={col_bobot_kriteria}{bobot_kriteria_start_row+i}")
        ws.cell(row=current_row+1, column=i+2).fill = fill_header
        
    ws.cell(row=current_row+1, column=1, value="Bobot Kriteria").font = font_bold
    ws.cell(row=current_row, column=n_k+2, value="SKOR AKHIR").font = font_bold
    
    # Isi Bobot Alternatif
    for i, a in enumerate(alternatif):
        ws.cell(row=current_row+2+i, column=1, value=a).font = font_bold
        for j, k in enumerate(kriteria):
            ref = bobot_alternatif_refs[k]
            ws.cell(row=current_row+2+i, column=j+2, value=f"={ref[2]}{ref[0]+i}")
            
        # Rumus Skor Akhir: SUMPRODUCT(baris bobot alternatif, baris bobot kriteria)
        start_col = get_column_letter(2)
        end_col = get_column_letter(n_k+1)
        ws.cell(row=current_row+2+i, column=n_k+2, value=f"=SUMPRODUCT({start_col}{current_row+2+i}:{end_col}{current_row+2+i}, {start_col}{current_row+1}:{end_col}{current_row+1})").fill = fill_result
        ws.cell(row=current_row+2+i, column=n_k+2).font = font_bold
        
    style_range(ws, current_row, 1, current_row+n_a+1, n_k+2)
    
    # Mengatur lebar kolom agar angka terlihat (menghindari ######)
    ws.column_dimensions['A'].width = 35
    for c in range(2, max(n_k, n_a) + 4):
        ws.column_dimensions[get_column_letter(c)].width = 18
    
    import io
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
